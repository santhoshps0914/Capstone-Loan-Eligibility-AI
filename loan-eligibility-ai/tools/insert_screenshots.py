"""Drop the captured screenshots into the 'Screenshots of Results' slide.

Each placeholder card keeps its existing fill, border, and caption styling; the
'[ SCREENSHOT PLACEHOLDER ]' line is removed, the caption is re-anchored to the
bottom of the card, and the PNG is centred in the space above it.

Run tools/capture_screenshots.py first.
"""
import os

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK = os.path.join(ROOT, "Loan_Eligibility_AI_Presentation.pptx")
SHOTS = os.path.join(ROOT, "assets", "screenshots")

PLACEHOLDER = "[ SCREENSHOT PLACEHOLDER ]"
SLIDE_TITLE = "Screenshots of Results"

# Card caption keyword -> screenshot file. Matched against the card's caption
# text so the mapping survives shapes being reordered in the XML.
MAPPING = [
    ("Application Form", "01_application_form.png"),
    ("Decision Panel", "02_decision_panel.png"),
    ("PDF Report", "03_pdf_report.png"),
    ("Audit Trail", "04_audit_trail.png"),
]

CAPTION_H = Inches(0.60)   # space reserved at the card bottom for the caption
PAD = Inches(0.08)         # inset between the image and the card border

NOTE_OLD = "Replace each frame"
NOTE_NEW = ("Captured from the running system (Streamlit :8501 → FastAPI :8000) — "
            "live decision, generated PDF, and the real append-only audit log.")


def find_slide(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and SLIDE_TITLE in shape.text_frame.text:
                return slide
    raise SystemExit(f"slide titled '{SLIDE_TITLE}' not found")


def strip_placeholder_line(shape):
    """Remove the placeholder paragraph (and the blank spacer) from a card."""
    tf = shape.text_frame
    for para in list(tf.paragraphs):
        text = "".join(r.text for r in para.runs)
        if PLACEHOLDER in text or (text.strip() == "" and len(tf.paragraphs) > 2):
            para._p.getparent().remove(para._p)


def fit(card, img_path):
    """Centre the image in the card above the caption, preserving aspect ratio."""
    avail_w = card.width - 2 * PAD
    avail_h = card.height - CAPTION_H - 2 * PAD
    with Image.open(img_path) as im:
        px_w, px_h = im.size
    scale = min(avail_w / px_w, avail_h / px_h)
    w, h = int(px_w * scale), int(px_h * scale)
    left = card.left + PAD + (avail_w - w) // 2
    top = card.top + PAD + (avail_h - h) // 2
    return Emu(left), Emu(top), Emu(w), Emu(h)


def main():
    prs = Presentation(DECK)
    slide = find_slide(prs)

    cards = {}
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text
        if PLACEHOLDER not in text:
            continue
        for keyword, filename in MAPPING:
            if keyword in text:
                cards[keyword] = (shape, filename)

    missing = [k for k, _ in MAPPING if k not in cards]
    if missing:
        raise SystemExit(f"no placeholder card found for: {missing}")

    for keyword, filename in MAPPING:
        card, _ = cards[keyword]
        path = os.path.join(SHOTS, filename)
        if not os.path.exists(path):
            raise SystemExit(f"missing screenshot: {path}")

        left, top, width, height = fit(card, path)
        strip_placeholder_line(card)
        # Caption sits in the reserved strip at the bottom of the card.
        card.text_frame.vertical_anchor = 4  # MSO_ANCHOR.BOTTOM
        slide.shapes.add_picture(path, left, top, width, height)
        print(f"inserted {filename} -> '{keyword}' card")

    for shape in slide.shapes:
        if shape.has_text_frame and NOTE_OLD in shape.text_frame.text:
            run = shape.text_frame.paragraphs[0].runs[0]
            run.text = NOTE_NEW
            print("updated the instruction note")

    prs.save(DECK)
    print(f"saved {DECK}")


if __name__ == "__main__":
    main()
